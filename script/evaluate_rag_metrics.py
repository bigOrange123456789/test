#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""只对已有 generations.jsonl 评分；不加载待测生成模型，也不打开图片。

示例：
    python script/evaluate_rag_metrics.py --generations A/generations.jsonl B/generations.jsonl
    python script/evaluate_rag_metrics.py --config script/evaluate_rag_metrics.json

评分依赖和答案解析复用 script/lib；此脚本不导入推理脚本。
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import io
import json
import logging
import math
import os
import re
import sys
import tempfile
import time
import unicodedata
from pathlib import Path
from typing import Any

import numpy as np
from tqdm import tqdm

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from script.lib.answer_metrics import question_type_for_sample, score_answer
from script.lib.atomic_factscore import AtomicFactScorer
from script.lib.evaluation_statistics import paired_summary, summarize_values
from script.lib.reference_judge import ReferenceJudge
from script.lib.token_budget import FACT_LENGTH_BUDGET_DEFAULTS, normalize_length_budget
from script.lib.text_metrics import compute_bleu4, compute_rouge_l

LOGGER = logging.getLogger("rag_eval_metrics")
DEFAULT_CONFIG = Path(__file__).with_name("evaluate_rag_metrics.json")
QUESTION_TYPES = ("open_ended", "closed_ended", "single_choice", "multiple_choice")
TYPE_LABELS = {"open_ended": "开放式问题", "closed_ended": "封闭式问题",
               "single_choice": "单选题", "multiple_choice": "多选题"}
METRICS = ("answer_accuracy", "choice_f1", "bleu4", "rouge_l", "semantic_score")
PRIMA_METRICS = ("factscore", "answer_accuracy", "explanation_rouge_l")
ALL_METRICS = tuple(dict.fromkeys((*PRIMA_METRICS, *METRICS)))
METRIC_LABELS = {"answer_accuracy": "Accuracy", "choice_f1": "Choice F1", "bleu4": "BLEU-4",
                 "rouge_l": "ROUGE-L", "semantic_score": "Semantic score",
                 "factscore": "FActScore", "explanation_rouge_l": "Explanation ROUGE-L"}
MODEL_PATHS = {
    "Qwen3-VL-2B-Instruct": PROJECT_ROOT / "Qwen3-VL-2B-Instruct",
    "DeepSeek-Model": PROJECT_ROOT / "DeepSeek-R1-Distill-Qwen-1.5B",
}
EVAL_DEFAULTS = {
    "judgeModel": "Qwen3-VL-2B-Instruct", "judgeModelPath": None,
    "judgeScope": "open_ended", "judgeMaxNewTokens": 256, "judgeRetries": 1,
    "bootstrapSamples": 1000,
    "prima": {
        "enabled": True, "separateQuestionTypes": True, "questionTypes": list(QUESTION_TYPES),
        "announceQuestionType": True,
        "factScore": True, "explanationRougeL": True, "factScoreMaxNewTokens": 1024,
        "factScoreBatchSize": 8, "factScoreLengthBudget": FACT_LENGTH_BUDGET_DEFAULTS,
    },
}


def _default_metric_selection(prima_enabled: bool, prima: dict) -> dict[str, bool]:
    """未配置逐项开关时保持旧版两套协议的行为。"""
    enabled = set(PRIMA_METRICS if prima_enabled else METRICS)
    return {name: name in enabled for name in ALL_METRICS}


def _normalize_metric_selection(raw: dict | None, evaluation: dict) -> dict[str, bool]:
    """规范化 7 个指标的开关；缺省时兼容原有 prima/legacy 行为。"""
    defaults = _default_metric_selection(evaluation["prima"]["enabled"], evaluation["prima"])
    if raw is None:
        return defaults
    if not isinstance(raw, dict):
        raise ValueError("evaluation.metrics 必须为 JSON 对象。")
    unknown = set(raw) - set(ALL_METRICS)
    if unknown:
        raise ValueError(f"evaluation.metrics 出现未知指标：{sorted(unknown)}")
    result = {**defaults, **raw}
    for name, value in result.items():
        if type(value) is not bool:
            raise ValueError(f"evaluation.metrics.{name} 必须为布尔值。")
    if not any(result.values()):
        raise ValueError("evaluation.metrics 至少要启用一个指标。")
    return result


def _canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def _fingerprint(value: Any) -> str:
    return hashlib.sha256(_canonical(value).encode("utf-8")).hexdigest()


def _read_jsonl(path: Path) -> list[dict]:
    rows = []
    with path.open("r", encoding="utf-8-sig") as stream:
        for number, line in enumerate(stream, 1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except ValueError as exc:
                raise ValueError(f"{path}:{number} JSON 解析失败：{exc}") from exc
            if not isinstance(row, dict):
                raise ValueError(f"{path}:{number} 必须为 JSON 对象。")
            rows.append(row)
    if not rows:
        raise ValueError(f"generations 文件为空：{path}")
    return rows


def _resolve_generations(value: str | Path, base: Path) -> list[Path]:
    path = Path(value).expanduser()
    path = (base / path).resolve() if not path.is_absolute() else path.resolve()
    if path.is_file():
        if path.name.casefold() != "generations.jsonl":
            raise ValueError(f"输入文件应为 generations.jsonl：{path}")
        return [path]
    if not path.is_dir():
        raise FileNotFoundError(f"找不到 generations 文件或目录：{path}")
    direct = path / "generations.jsonl"
    if direct.is_file():
        return [direct]
    found = sorted(path.rglob("generations.jsonl"))
    if not found:
        raise FileNotFoundError(f"目录中没有 generations.jsonl：{path}")
    return found


def _safe_label(value: str) -> str:
    # Windows 会把带扩展名的设备名也视为保留名，例如 CON.txt。
    device_stem = value.split(".", 1)[0].rstrip(" .").upper() if isinstance(value, str) else ""
    if (not isinstance(value, str) or not value.strip() or value != value.strip()
            or value in {".", ".."} or value.endswith(".")
            or not re.fullmatch(r"[\w .\-\u4e00-\u9fff]{1,100}", value)
            or device_stem in {"CON", "PRN", "AUX", "NUL"}
            or re.fullmatch(r"(?:COM|LPT)(?:[1-9]|[¹²³])", device_stem, re.I)):
        raise ValueError(f"标签不能安全用作目录名：{value!r}")
    return value


def _label_directory(root: Path, label: str) -> Path:
    """解析并确认标签目录始终位于本次评分输出根目录内。"""
    root = root.resolve()
    directory = (root / _safe_label(label)).resolve()
    if directory == root or root not in directory.parents:
        raise ValueError(f"评分标签目录越出输出根目录：{directory}")
    return directory


def _read_sidecar_evaluation(path: Path) -> dict | None:
    sidecar = path.parent / "summary.json"
    if not sidecar.is_file():
        return None
    try:
        summary = json.loads(sidecar.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError) as exc:
        raise ValueError(f"无法读取 generations 对应的 summary.json：{sidecar}: {exc}") from exc
    value = summary.get("evaluation") if isinstance(summary, dict) else None
    return value if isinstance(value, dict) else None


def _normalize_evaluation(raw: dict | None, config_dir: Path) -> dict:
    raw = raw or {}
    if not isinstance(raw, dict):
        raise ValueError("evaluation 必须为 JSON 对象。")
    # metrics 是评分阶段的逐项开关；允许它出现在 evaluation 中，
    # 但在下面单独规范化，避免把原始生成阶段的 sidecar 配置误当成评分配置。
    metrics_raw = raw.get("metrics")
    protocol_raw = {key: value for key, value in raw.items() if key != "metrics"}
    unknown = set(protocol_raw) - set(EVAL_DEFAULTS)
    if unknown:
        raise ValueError(f"evaluation 出现未知配置项：{sorted(unknown)}")
    result = {**EVAL_DEFAULTS, **protocol_raw}
    prima_raw = protocol_raw.get("prima", {})
    if not isinstance(prima_raw, dict) or set(prima_raw) - (set(EVAL_DEFAULTS["prima"]) | {"announceQuestionType"}):
        raise ValueError(f"evaluation.prima 配置无效；可选字段为 {list(EVAL_DEFAULTS['prima'])}")
    prima = {**EVAL_DEFAULTS["prima"], **prima_raw}
    for field in ("enabled", "separateQuestionTypes", "factScore", "explanationRougeL"):
        if type(prima[field]) is not bool:
            raise ValueError(f"evaluation.prima.{field} 必须为布尔值。")
    types = prima["questionTypes"]
    if (not isinstance(types, list) or not types or len(set(types)) != len(types)
            or any(item not in QUESTION_TYPES for item in types)):
        raise ValueError(f"evaluation.prima.questionTypes 必须是有效题型列表：{QUESTION_TYPES}")
    for field, low, high in (("factScoreMaxNewTokens", 128, 8192), ("factScoreBatchSize", 1, 32)):
        if type(prima[field]) is not int or not low <= prima[field] <= high:
            raise ValueError(f"evaluation.prima.{field} 范围为 {low}～{high}。")
    budget = prima["factScoreLengthBudget"]
    if not isinstance(budget, dict):
        raise ValueError("evaluation.prima.factScoreLengthBudget 必须为对象。")
    prima["factScoreLengthBudget"] = normalize_length_budget(
        budget, "evaluation.prima.factScoreLengthBudget", defaults=FACT_LENGTH_BUDGET_DEFAULTS)
    result["prima"] = prima
    if result["judgeModel"] not in MODEL_PATHS:
        raise ValueError(f"judgeModel 仅支持：{', '.join(MODEL_PATHS)}")
    if result["judgeScope"] not in {"none", "open_ended", "all"}:
        raise ValueError("judgeScope 可选 none、open_ended、all。")
    for field, low, high in (("judgeMaxNewTokens", 32, 2048), ("judgeRetries", 0, 1),
                             ("bootstrapSamples", 100, 10000)):
        if type(result[field]) is not int or not low <= result[field] <= high:
            raise ValueError(f"evaluation.{field} 范围为 {low}～{high}。")
    judge_path = result["judgeModelPath"]
    if judge_path is None:
        result["judgeModelPath"] = str(MODEL_PATHS[result["judgeModel"]].resolve())
    elif isinstance(judge_path, str) and judge_path.strip():
        candidate = Path(judge_path).expanduser()
        result["judgeModelPath"] = str((config_dir / candidate).resolve() if not candidate.is_absolute() else candidate.resolve())
    else:
        raise ValueError("evaluation.judgeModelPath 必须为路径字符串或 null。")
    result["metrics"] = _normalize_metric_selection(metrics_raw, result)
    result["_metrics_explicit"] = metrics_raw is not None
    return result


class _TextJudgeModel:
    """按需加载固定裁判模型，仅提供文本接口；绝不接收图片。"""

    def __init__(self, model_type: str, model_path: str, *, max_input_tokens: int = 16384):
        import torch

        self.torch = torch
        self.model_type, self.model_path = model_type, model_path
        self.max_input_tokens = max_input_tokens
        self.model = self.tokenizer = self.processor = None
        self.last_generation_info = {}
        self.system_prompt = "你是严格的医学事实核验助手。只依据给定的参考资料评分，不使用外部知识。"
        try:
            from script.lib.prompt_config import prompt
            self.system_prompt = prompt("evaluation.fact_system", default=self.system_prompt)
        except (FileNotFoundError, KeyError):
            pass

    def _load(self):
        if self.model is not None:
            return
        from transformers import AutoModelForCausalLM, AutoProcessor, AutoTokenizer
        device = "cuda" if self.torch.cuda.is_available() else "cpu"
        dtype = (self.torch.bfloat16 if device == "cuda" and self.torch.cuda.is_bf16_supported()
                 else self.torch.float16 if device == "cuda" else self.torch.float32)
        if self.model_type == "DeepSeek-Model":
            self.tokenizer = AutoTokenizer.from_pretrained(self.model_path, use_fast=True, local_files_only=True)
            if self.tokenizer.pad_token_id is None:
                self.tokenizer.pad_token = self.tokenizer.eos_token
            self.model = AutoModelForCausalLM.from_pretrained(
                self.model_path, dtype=dtype, low_cpu_mem_usage=True, local_files_only=True).to(device).eval()
        else:
            try:
                from transformers import AutoModelForVision2Seq
                self.model = AutoModelForVision2Seq.from_pretrained(
                    self.model_path, dtype=dtype, local_files_only=True)
            except (ImportError, ValueError):
                from transformers import AutoModelForImageTextToText
                self.model = AutoModelForImageTextToText.from_pretrained(
                    self.model_path, dtype=dtype, local_files_only=True)
            self.model = self.model.to(device).eval()
            self.processor = AutoProcessor.from_pretrained(self.model_path, local_files_only=True)
            self.tokenizer = self.processor.tokenizer
        self.device = device

    def count_tokens(self, text: str) -> int:
        self._load()
        return len(self.tokenizer.encode(text, add_special_tokens=False))

    def text(self, prompt_text: str, max_new_tokens: int = 256) -> str:
        self._load()
        if self.model_type == "DeepSeek-Model":
            messages = [{"role": "system", "content": self.system_prompt},
                        {"role": "user", "content": prompt_text}]
            rendered = self.tokenizer.apply_chat_template(
                messages + [{"role": "assistant", "content": ""}],
                tokenize=False, add_generation_prompt=False)
            eos = self.tokenizer.eos_token
            if eos and rendered.endswith(eos):
                rendered = rendered[:-len(eos)]
            inputs = self.tokenizer(rendered, return_tensors="pt", truncation=False, add_special_tokens=False)
            decode = lambda ids: self.tokenizer.decode(ids, skip_special_tokens=True,
                                                       clean_up_tokenization_spaces=False)
        else:
            messages = [{"role": "system", "content": self.system_prompt},
                        {"role": "user", "content": prompt_text}]
            rendered = self.processor.apply_chat_template(messages, tokenize=False, add_generation_prompt=True)
            inputs = self.processor(text=[rendered], images=None, videos=None, return_tensors="pt")
            decode = lambda ids: self.processor.batch_decode(
                ids, skip_special_tokens=True, clean_up_tokenization_spaces=False)[0]
        input_length = int(inputs["input_ids"].shape[1])
        if input_length > self.max_input_tokens:
            raise ValueError(f"裁判输入 {input_length} tokens 超过 max_input_tokens={self.max_input_tokens}。")
        inputs = inputs.to(self.device)
        self.model.config.use_cache = True
        started = time.perf_counter()
        with self.torch.inference_mode():
            output = self.model.generate(**inputs, max_new_tokens=max_new_tokens, do_sample=False,
                                         num_beams=1, pad_token_id=self.tokenizer.pad_token_id,
                                         eos_token_id=self.tokenizer.eos_token_id)
        new_ids = output[0, input_length:]
        eos_ids = self.tokenizer.eos_token_id
        eos_ids = set(eos_ids) if isinstance(eos_ids, (list, tuple, set)) else {eos_ids}
        ended = bool(new_ids.numel() and int(new_ids[-1]) in eos_ids)
        self.last_generation_info = {"input_tokens": input_length, "output_tokens": int(new_ids.shape[0]),
                                     "hit_token_limit": int(new_ids.shape[0]) >= max_new_tokens and not ended,
                                     "max_new_tokens": max_new_tokens,
                                     "elapsed_seconds": time.perf_counter() - started}
        return decode(new_ids).strip()

    def close(self):
        self.model = self.tokenizer = self.processor = None
        import gc
        gc.collect()
        if self.torch.cuda.is_available():
            self.torch.cuda.empty_cache()


def _metric_names(evaluation: dict) -> tuple[str, ...]:
    selection = evaluation.get("metrics")
    if isinstance(selection, dict):
        return tuple(name for name in ALL_METRICS if selection.get(name, False))
    return PRIMA_METRICS if evaluation["prima"]["enabled"] else METRICS


def _metric_applies(row: dict, metric: str) -> bool:
    kind = row["answer_details"]["question_type"]
    if metric == "answer_accuracy":
        return kind in {"single_choice", "multiple_choice", "closed_ended"}
    if metric == "choice_f1":
        return kind in {"single_choice", "multiple_choice"}
    if metric == "factscore":
        return row.get("factscore_required", False)
    if metric == "semantic_score":
        return row.get("semantic_required", False)
    if metric == "explanation_rouge_l":
        return row["explanation_enabled"]
    return True


def _summarize(rows: list[dict], evaluation: dict, seed: int) -> dict:
    result = {}
    prima = evaluation["prima"]
    for metric in _metric_names(evaluation):
        selected = [row for row in rows if _metric_applies(row, metric)]
        result[metric] = summarize_values([row.get(metric) for row in selected], seed=seed,
                                          bootstrap_samples=evaluation["bootstrapSamples"],
                                          binary=metric == "answer_accuracy")
    return result


def _update_summary(summary: dict, rows: list[dict], seed: int) -> None:
    evaluation = summary["evaluation"]
    summary["metrics"] = _summarize(rows, evaluation, seed)
    types = QUESTION_TYPES if evaluation["prima"]["enabled"] else sorted(
        {row["answer_details"]["question_type"] for row in rows})
    summary["by_question_type"] = {
        kind: {"num_samples": sum(row["answer_details"]["question_type"] == kind for row in rows),
               "metrics": _summarize([row for row in rows if row["answer_details"]["question_type"] == kind],
                                     evaluation, seed)}
        for kind in types
    }
    summary["judge_failures"] = sum(
        any(details.get("status") == "judge_failed" for details in (
            row.get("factscore_details", {}), row.get("semantic_details", {})))
        for row in rows)
    summary["judge_required"] = sum(row["judge_required"] for row in rows)
    summary["possible_truncations"] = sum(bool(row.get("generation_info", {}).get("hit_token_limit")) for row in rows)
    summary["review_count"] = sum(_needs_review(row) for row in rows)
    summary["scoring_status"] = "needs_review" if summary["judge_failures"] else "complete"


def _needs_review(row: dict) -> bool:
    judge_failed = any(details.get("status") == "judge_failed" for details in (
        row.get("factscore_details", {}), row.get("semantic_details", {})))
    return bool(judge_failed
                or row.get("answer_details", {}).get("parse_status") in {"prediction_unparseable", "invalid_reference"}
                or row.get("generation_info", {}).get("hit_token_limit")
                or (row.get("explanation_enabled") and row["answer_details"].get("explanation_status") != "ok"))


def _load_dataset_fallback(dataset_path: str | None, paths: list[Path], records: list[dict]) -> list[dict]:
    needs_answer_fields = any("answer" not in row or ("options" not in row and "Options: " not in row.get("question", ""))
                             for row in records)
    if not needs_answer_fields:
        return records
    if dataset_path is None:
        # Old generations may have the serialized answer in reference; score_answer handles JSON text.
        return records
    source = Path(dataset_path).expanduser().resolve()
    identifiers = list(dict.fromkeys(str(row["id"]) for row in records))
    from script.lib.evaluation_data import _load_jsonl, _load_selected_mira
    if source.is_dir():
        source_records = _load_selected_mira(source, identifiers)
    else:
        by_id = {str(row["id"]): row for row in _load_jsonl(source)}
        missing = [identifier for identifier in identifiers if identifier not in by_id]
        if missing:
            raise ValueError(f"回源数据缺少 {len(missing)} 个 ID，例如：{missing[:5]}")
        source_records = [by_id[identifier] for identifier in identifiers]
    by_id = {str(row["id"]): row for row in source_records}
    return [{**by_id.get(str(row["id"]), {}), **row,
             **{key: by_id[str(row["id"])][key] for key in ("answer", "options", "question_type")
                if key not in row and key in by_id.get(str(row["id"]), {})}}
            for row in records]


def _make_rows(records: list[dict], evaluation: dict, label: str) -> list[dict]:
    prima = evaluation["prima"]
    selected_metrics = evaluation["metrics"]
    rows = []
    seen = set()
    for source in records:
        if "id" not in source or str(source["id"]).strip() == "":
            raise ValueError(f"{label}: generations 记录缺少 id。")
        key = str(source["id"]).strip()
        if key in seen:
            raise ValueError(f"{label}: generations 包含重复 id：{key}")
        seen.add(key)
        question = source.get("question")
        prediction = source.get("prediction")
        reference = source.get("reference", source.get("answer"))
        if not isinstance(question, str) or not question.strip() or not isinstance(prediction, str):
            raise ValueError(f"{label}/{key}: question 与 prediction 必须为文本。")
        if reference is None:
            raise ValueError(f"{label}/{key}: 缺少 reference/answer。")
        sample = {**source, "reference": reference}
        sample.setdefault("answer", reference)
        answers = score_answer(sample, prediction)
        is_open = answers["question_type"] == "open_ended"
        factscore_required = bool(selected_metrics.get("factscore") and is_open
                                  and (not prima["enabled"] or "open_ended" in prima["questionTypes"])
                                  and (evaluation.get("_metrics_explicit") or prima["factScore"]))
        semantic_required = bool(selected_metrics.get("semantic_score") and
                                 (evaluation["judgeScope"] == "all" or
                                  evaluation["judgeScope"] == answers["question_type"] or
                                  (evaluation["judgeScope"] == "open_ended" and is_open)))
        judge_required = factscore_required or semantic_required
        judge_metric = ("both" if factscore_required and semantic_required
                        else "factscore" if factscore_required else "semantic_score")
        explanation_enabled = bool(selected_metrics.get("explanation_rouge_l") and
                                   (evaluation.get("_metrics_explicit") or prima["explanationRougeL"]))
        explanation_score = None
        if explanation_enabled and answers["explanation_eligible"]:
            explanation_score = compute_rouge_l(answers["prediction_explanation"], answers["reference_explanation"])
        normalized_prediction = answers["normalized_prediction"]
        normalized_reference = answers["normalized_reference"]
        row = {**sample, "prediction": prediction, "question_type": answers["question_type"],
               "normalized_prediction": normalized_prediction, "normalized_reference": normalized_reference,
               "answer_details": answers, "answer_accuracy": answers["answer_accuracy"],
               "choice_f1": answers["choice_f1"], "bleu4": compute_bleu4(normalized_prediction, normalized_reference),
               "rouge_l": compute_rouge_l(normalized_prediction, normalized_reference),
               "semantic_score": None, "factscore": None, "explanation_rouge_l": explanation_score,
               "explanation_enabled": explanation_enabled, "judge_required": judge_required,
               "factscore_required": factscore_required, "semantic_required": semantic_required,
               "judge_metric": judge_metric,
               "semantic_details": {"score": None, "status": "pending" if judge_required and judge_metric == "semantic_score" else "not_applicable"},
               "factscore_details": {"score": None, "status": "pending" if judge_required and judge_metric == "factscore" else "not_applicable"},
               "retrieved_evidence": source.get("retrieved_evidence", []),
               "retrieved_count": len(source.get("retrieved_evidence", [])),
               "config_name": source.get("config_name"),
               "generation_fingerprint": source.get("generation_fingerprint"),
               "generation_from_cache": False, "generation_info": source.get("generation_info", {})}
        rows.append(row)
    return rows


def _judge_rows(rows: list[dict], evaluation: dict, llm, cache_dir: Path, identity: dict) -> None:
    prima = evaluation["prima"]
    fact_rows = [row for row in rows if row.get("factscore_required")]
    semantic_rows = [row for row in rows if row.get("semantic_required")]
    if fact_rows:
        scorer = AtomicFactScorer(llm, cache_dir, identity,
                                  max_new_tokens=prima["factScoreMaxNewTokens"],
                                  batch_size=prima["factScoreBatchSize"], retries=evaluation["judgeRetries"],
                                  length_budget=prima["factScoreLengthBudget"])
        for row in tqdm(fact_rows, desc="FActScore"):
            answers = row["answer_details"]
            sample = {**row, "evaluation_reference": answers["semantic_reference"]}
            details = scorer.score(sample, answers["semantic_prediction"])
            row["factscore_details"] = details
            row["factscore"] = details["score"]
    if semantic_rows:
        scorer = ReferenceJudge(llm, cache_dir, identity,
                                max_new_tokens=evaluation["judgeMaxNewTokens"], retries=evaluation["judgeRetries"])
        for row in tqdm(semantic_rows, desc="语义裁判"):
            answers = row["answer_details"]
            sample = {**row, "evaluation_reference": answers["semantic_reference"]}
            details = scorer.score(sample, answers["semantic_prediction"])
            row["semantic_details"] = details
            row["semantic_score"] = details["score"]


def _paired_metrics(after: dict, before: dict, evaluation: dict, seed: int) -> dict:
    if after.keys() != before.keys():
        raise ValueError("配对差值要求两组 generations 的 ID 集合完全一致。")
    result = {}
    for metric in _metric_names(evaluation):
        keys = [key for key in sorted(after) if _metric_applies(after[key], metric) and _metric_applies(before[key], metric)]
        result[metric] = paired_summary([after[key][metric] for key in keys], [before[key][metric] for key in keys],
                                        seed=seed, bootstrap_samples=evaluation["bootstrapSamples"])
    return result


def _summary(label: str, generation_path: Path, source_records: list[dict], rows: list[dict], evaluation: dict,
             seed: int, elapsed: float, sidecar: dict | None = None) -> dict:
    test_fingerprint = _fingerprint([{key: row.get(key) for key in
                                      ("id", "question", "reference", "answer", "options", "images", "question_type")}
                                     for row in source_records])
    sidecar = sidecar or {}
    persisted_evaluation = {key: value for key, value in evaluation.items() if key != "_metrics_explicit"}
    return {"name": label, "config_name": rows[0].get("config_name"), "generations_path": str(generation_path),
            "num_samples": len(rows), "test_fingerprint": test_fingerprint,
            "dataset_fingerprint": sidecar.get("dataset_fingerprint", test_fingerprint),
            "generation_fingerprint": rows[0].get("generation_fingerprint"),
            "evaluation": persisted_evaluation, "metric_scale": "0-1", "std_ddof": 0,
            "aggregation": "macro average over test samples", "protocol": "prima" if evaluation["prima"]["enabled"] else "reference_semantic",
            "model": sidecar.get("model"), "model_path": sidecar.get("model_path"),
            "lora_path": sidecar.get("lora_path"), "use_rag": sidecar.get("use_rag"),
            "top_k": sidecar.get("top_k"), "chroma_db_dir": sidecar.get("chroma_db_dir"),
            "chroma_collection": sidecar.get("chroma_collection"),
            "chroma_id_prefix": sidecar.get("chroma_id_prefix"),
            "scoring_version": sidecar.get("scoring_version"),
            "bleu_smoothing": sidecar.get("bleu_smoothing"),
            "generation_settings": sidecar.get("generation_settings"), "seed": seed,
            "generation_seconds_excluding_load": sum(float(row.get("generation_info", {}).get("elapsed_seconds", 0) or 0) for row in rows),
            "generation_input_mode": sidecar.get("generation_input_mode", "from_generations"), "scoring_seconds": elapsed}


def _save_one(label: str, generation_path: Path, records: list[dict], evaluation: dict,
              seed: int, judge, judge_identity: dict | None, cache_dir: Path,
              sidecar: dict | None = None) -> tuple[dict, list[dict]]:
    started = time.perf_counter()
    records = _load_dataset_fallback(None, [generation_path], records)
    rows = _make_rows(records, evaluation, label)
    if any(row["judge_required"] for row in rows):
        if judge is None:
            raise RuntimeError("当前配置需要裁判，但裁判模型未初始化。")
        _judge_rows(rows, evaluation, judge, cache_dir, judge_identity or {})
    summary = _summary(label, generation_path, records, rows, evaluation, seed,
                       time.perf_counter() - started, sidecar)
    _update_summary(summary, rows, seed)
    return summary, rows


def _summary_display(metrics: dict) -> str:
    value = metrics.get("mean")
    return "—" if value is None else f"{value * 100:.2f} ± {metrics['std'] * 100:.2f} (n={metrics['n']})"


CONSOLE_METRIC_LABELS = {
    "factscore": "FActScore (%)",
    "answer_accuracy": "Accuracy (%)",
    "explanation_rouge_l": "解释ROUGE-L (%)",
    "choice_f1": "Choice F1 (%)",
    "bleu4": "BLEU-4 (%)",
    "rouge_l": "ROUGE-L (%)",
    "semantic_score": "语义评分 (%)",
}


def _display_width(value: str) -> int:
    """估算 CMD 中中英文混排文本占用的显示宽度。"""
    return sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in value)


def _pad_console_cell(value: str, width: int) -> str:
    return value + " " * max(0, width - _display_width(value))


def _console_table(summaries: list[dict], evaluation: dict) -> str:
    """生成简洁的题型分组汇总表，供 CMD 查看；详细结果仍保存为 JSON/Markdown。"""
    metric_names = _metric_names(evaluation)
    headers = ["任务", "题型", "N", *(CONSOLE_METRIC_LABELS[name] for name in metric_names)]
    show_explanation_count = "explanation_rouge_l" in metric_names
    if show_explanation_count:
        headers.append("有效解释/题数")
    rows: list[list[str]] = []
    for summary in summaries:
        groups = summary.get("by_question_type", {}) if evaluation["prima"]["separateQuestionTypes"] else {"all": summary}
        if not groups:
            groups = {"all": summary}
        for kind, group in groups.items():
            metrics = group.get("metrics", {})
            values = [_summary_display(metrics.get(name, {})) for name in metric_names]
            if show_explanation_count:
                explanation = metrics.get("explanation_rouge_l", {})
                valid = explanation.get("n", 0) if isinstance(explanation, dict) else 0
                values.append(f"{valid}/{group.get('num_samples', 0)}")
            rows.append([summary.get("name", ""), TYPE_LABELS.get(kind, "合并"),
                         str(group.get("num_samples", 0)), *values])
    if not rows:
        return ""
    widths = [max(_display_width(headers[index]), *(_display_width(row[index]) for row in rows))
              for index in range(len(headers))]
    separator = "+".join("-" * width for width in widths)
    lines = [" | ".join(_pad_console_cell(value, widths[index]) for index, value in enumerate(headers)),
             separator]
    lines.extend(" | ".join(_pad_console_cell(value, widths[index]) for index, value in enumerate(row))
                 for row in rows)
    return "\n".join(lines)


def _write_reports(root: Path, summaries: list[dict], by_label: dict[str, dict], evaluation: dict, seed: int) -> dict:
    paired = []
    names = list(by_label)
    # 默认报告 LoRA 与其原版的同题差值；仅在两组 ID 集合完全相同时计算。
    summary_by_name = {summary["name"]: summary for summary in summaries}
    for fine_name in names:
        fine_summary = summary_by_name[fine_name]
        fine_lora = bool(fine_summary.get("lora_path") or re.search(r"lora|微调", fine_name, re.I))
        if not fine_lora:
            continue
        base_candidates = [name for name in names if name != fine_name and not (
            summary_by_name[name].get("lora_path") or re.search(r"lora|微调", name, re.I))]
        compatible_candidates = []
        for base_name in base_candidates:
            base_summary = summary_by_name[base_name]
            compatible_fields = ["model_path", "use_rag", "top_k", "test_fingerprint", "dataset_fingerprint",
                                 "evaluation", "scoring_version", "bleu_smoothing", "generation_input_mode",
                                 "generation_settings"]
            compatible_fields.append("model")
            if any(fine_summary.get(field) != base_summary.get(field) for field in compatible_fields):
                continue
            if fine_summary.get("use_rag") and any(
                    fine_summary.get(field) != base_summary.get(field)
                    for field in ("chroma_db_dir", "chroma_collection", "chroma_id_prefix")):
                continue
            compatible_candidates.append(base_name)
        if not fine_summary.get("model") and len(compatible_candidates) > 1:
            # 未提供 sidecar 模型身份时，不从多个候选原版中猜测配对关系。
            compatible_candidates = []
        for base_name in compatible_candidates:
            base_summary = summary_by_name[base_name]
            after, before = by_label[fine_name], by_label[base_name]
            after_map = {str(row["id"]): row for row in after}
            before_map = {str(row["id"]): row for row in before}
            if after_map.keys() != before_map.keys():
                LOGGER.warning("未计算 %s 相对 %s 的配对差值：ID 集合不一致（%d vs %d）。",
                               fine_name, base_name, len(after_map), len(before_map))
                continue
            delta = {"name": fine_name, "baseline": base_name, "num_samples": len(after_map),
                     "metrics": _paired_metrics(after_map, before_map, evaluation, seed)}
            if evaluation["prima"]["enabled"]:
                delta["by_question_type"] = {
                    kind: _paired_metrics(
                        {key: row for key, row in after_map.items() if row["answer_details"]["question_type"] == kind},
                        {key: row for key, row in before_map.items() if row["answer_details"]["question_type"] == kind},
                        evaluation, seed)
                    for kind in QUESTION_TYPES}
            paired.append(delta)
    same_data = len({summary["test_fingerprint"] for summary in summaries}) <= 1
    result = {"schema_version": "rag-metrics-v1", "status": "complete", "protocol": "prima" if evaluation["prima"]["enabled"] else "reference_semantic",
              "same_test_data": same_data, "metrics_enabled": list(_metric_names(evaluation)),
              "summaries": summaries, "paired_deltas": paired,
              "note": "多选题按所选选项集合精确匹配；缺失评分记 null 并列入人工复核，不以零分替代。"}
    question_rows = []
    report_metrics = _metric_names(evaluation)
    for summary in summaries:
        groups = summary.get("by_question_type", {}) if evaluation["prima"]["separateQuestionTypes"] else {"all": summary}
        if not groups:
            groups = {"all": summary}
        for kind, group in groups.items():
            metrics = group["metrics"]
            row = {"name": summary["name"], "question_type": kind, "num_samples": group["num_samples"],
                   "metrics": metrics}
            question_rows.append(row)
    result["question_type_comparison"] = question_rows
    return result


EXCEL_CHUNK_SIZE = 15000
EXCEL_CELL_LIMIT = 30000


def _json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


def _excel_value(value: Any, *, truncate: bool = True) -> Any:
    if isinstance(value, (dict, list, tuple)):
        value = _json_text(value)
    if isinstance(value, str):
        if value.lstrip().startswith(("=", "+", "-", "@")):
            value = "'" + value
        if truncate and len(value) > EXCEL_CELL_LIMIT:
            value = value[:EXCEL_CELL_LIMIT - 40] + "…（内容过长，完整内容见原始记录页）"
    return value


def _style_sheet(sheet, headers: list[str], rows: list[list[Any]]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet.append(headers)
    for row in rows:
        sheet.append([_excel_value(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315B64")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for column_index, header in enumerate(headers, 1):
        samples = [str(header)]
        samples.extend(str(row[column_index - 1]) for row in rows[:300]
                       if column_index - 1 < len(row) and row[column_index - 1] is not None)
        width = min(60, max(12, max((min(len(value), 60) for value in samples), default=12) + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    if len(headers) == 1 and rows:
        sheet.column_dimensions["A"].width = 100


def _add_sheet(workbook, title: str, headers: list[str], rows: list[list[Any]]):
    sheet = workbook.create_sheet(title)
    _style_sheet(sheet, headers, rows)
    return sheet


def _percent(value: Any) -> float | None:
    return None if value is None else float(value) * 100


def _review_reasons(row: dict) -> str:
    reasons = []
    answers = row.get("answer_details", {})
    if answers.get("parse_status") in {"prediction_unparseable", "invalid_reference"}:
        reasons.append(f"答案解析：{answers['parse_status']}")
    if row.get("generation_info", {}).get("hit_token_limit"):
        reasons.append("生成回答可能达到 token 上限")
    if row.get("explanation_enabled") and answers.get("explanation_status") != "ok":
        reasons.append(f"解释文本：{answers.get('explanation_status', '缺失')}")
    for label, key in (("FActScore", "factscore_details"), ("语义评分", "semantic_details")):
        if row.get(key, {}).get("status") == "judge_failed":
            reasons.append(f"{label}裁判失败")
    return "；".join(reasons)


def _metric_summary_rows(summaries: list[dict], evaluation: dict) -> list[list[Any]]:
    rows = []
    for summary in summaries:
        groups = summary.get("by_question_type", {}) if evaluation["prima"]["separateQuestionTypes"] else {"all": summary}
        for kind, group in (groups or {"all": summary}).items():
            for metric in _metric_names(evaluation):
                stats = group.get("metrics", {}).get(metric, {})
                ci95 = stats.get("ci95", [None, None])
                rows.append([
                    summary.get("name"), TYPE_LABELS.get(kind, "合并"), group.get("num_samples", 0),
                    METRIC_LABELS[metric], stats.get("n", 0), stats.get("total", group.get("num_samples", 0)),
                    _percent(stats.get("coverage")), _percent(stats.get("mean")), _percent(stats.get("std")),
                    _percent(ci95[0] if ci95 else None), _percent(ci95[1] if len(ci95) > 1 else None),
                    stats.get("ci_method"),
                ])
    return rows


def _detail_headers(evaluation: dict) -> list[str]:
    return [
        "评估名称", "题目ID", "题型", "问题", "标准答案", "参考文本", "生成回答",
        "归一化回答", "归一化参考",
        *(METRIC_LABELS[name] + " (%)" for name in ALL_METRICS),
        "答案解析状态", "答案解析详情", "FActScore状态", "FActScore原因",
        "语义评分状态", "语义评分原因", "解释状态", "生成信息", "检索证据数",
        "是否需复核", "复核原因",
    ]


def _detail_row(label: str, row: dict) -> list[Any]:
    answers = row.get("answer_details", {})
    fact = row.get("factscore_details", {})
    semantic = row.get("semantic_details", {})
    reasons = _review_reasons(row)
    return [
        label, row.get("id"), TYPE_LABELS.get(row.get("question_type"), row.get("question_type")),
        row.get("question"), row.get("answer", row.get("reference")), row.get("reference"),
        row.get("prediction"), row.get("normalized_prediction"), row.get("normalized_reference"),
        *(_percent(row.get(metric)) for metric in ALL_METRICS),
        answers.get("parse_status"), _json_text(answers), fact.get("status"), fact.get("reason"),
        semantic.get("status"), semantic.get("reason"), answers.get("explanation_status"),
        _json_text(row.get("generation_info", {})), row.get("retrieved_count", 0),
        bool(reasons), reasons,
    ]


def _append_chunked_json(rows: list[list[Any]], category: str, object_name: str, field: str, value: Any) -> None:
    content = _json_text(value)
    chunks = [content[index:index + EXCEL_CHUNK_SIZE] for index in range(0, len(content), EXCEL_CHUNK_SIZE)] or [""]
    for index, chunk in enumerate(chunks, 1):
        rows.append([category, object_name, field, chunk, index, len(chunks)])


def _restore_judge_cache(workbook_path: Path, cache_dir: Path) -> int:
    if not workbook_path.is_file():
        return 0
    from openpyxl import load_workbook

    workbook = None
    restored = 0
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if "_judge_cache" not in workbook.sheetnames:
            return 0
        grouped: dict[str, list[tuple[int, str]]] = {}
        for relative, part, content in workbook["_judge_cache"].iter_rows(min_row=2, values_only=True):
            if not isinstance(relative, str) or type(part) is not int or not isinstance(content, str):
                continue
            grouped.setdefault(relative, []).append((part, content))
        cache_root = cache_dir.resolve()
        for relative, chunks in grouped.items():
            target = (cache_root / relative).resolve()
            if target == cache_root or cache_root not in target.parents:
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text("".join(piece for _, piece in sorted(chunks)), encoding="utf-8")
            restored += 1
    except Exception as exc:
        LOGGER.warning("未能从现有工作簿恢复裁判缓存，将重新评分：%s", exc)
    finally:
        if workbook is not None:
            workbook.close()
    return restored


def _judge_cache_rows(cache_dir: Path) -> list[list[Any]]:
    rows = []
    if not cache_dir.is_dir():
        return rows
    for path in sorted(cache_dir.rglob("*.json")):
        relative = path.relative_to(cache_dir).as_posix()
        content = path.read_text(encoding="utf-8")
        chunks = [content[index:index + EXCEL_CHUNK_SIZE] for index in range(0, len(content), EXCEL_CHUNK_SIZE)] or [""]
        rows.extend([[relative, index, chunk] for index, chunk in enumerate(chunks, 1)])
    return rows


def _write_workbook(workbook_path: Path, task: dict, summaries: list[dict], by_label: dict[str, list[dict]],
                    evaluation: dict, seed: int, report: dict, manifest: dict, log_text: str,
                    cache_dir: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_sheet = _add_sheet(
        workbook, "指标汇总",
        ["评估名称", "题型", "题数", "指标", "有效评分数", "适用题数", "覆盖率 (%)",
         "平均值 (%)", "标准差（百分点）", "95% CI下限 (%)", "95% CI上限 (%)", "区间方法"],
        _metric_summary_rows(summaries, evaluation),
    )
    for row in summary_sheet.iter_rows(min_row=2, min_col=7, max_col=11):
        for cell in row:
            cell.number_format = "0.00"

    detail_rows = []
    review_rows = []
    evidence_rows = []
    raw_rows = []
    for summary in summaries:
        label = summary["name"]
        for row in by_label[label]:
            detail_rows.append(_detail_row(label, row))
            reasons = _review_reasons(row)
            if reasons:
                review_rows.append([
                    label, row.get("id"), TYPE_LABELS.get(row.get("question_type"), row.get("question_type")),
                    row.get("question"), row.get("reference"), row.get("prediction"), reasons,
                ])
            evidence = row.get("retrieved_evidence", [])
            if isinstance(evidence, dict):
                evidence = [evidence]
            if isinstance(evidence, list):
                for index, item in enumerate(evidence, 1):
                    item = item if isinstance(item, dict) else {"内容": item}
                    evidence_rows.append([
                        label, row.get("id"), index,
                        item.get("text", item.get("document", item.get("content"))),
                        item.get("score", item.get("distance")), _json_text(item),
                    ])
            raw = _json_text(row)
            chunks = [raw[index:index + EXCEL_CHUNK_SIZE] for index in range(0, len(raw), EXCEL_CHUNK_SIZE)] or [""]
            raw_rows.extend([[label, row.get("id"), index, len(chunks), chunk]
                             for index, chunk in enumerate(chunks, 1)])

    _add_sheet(workbook, "逐题明细", _detail_headers(evaluation), detail_rows)
    _add_sheet(workbook, "复核清单",
               ["评估名称", "题目ID", "题型", "问题", "参考文本", "生成回答", "复核原因"],
               review_rows)

    paired_rows = []
    for delta in report.get("paired_deltas", []):
        groups = [("总体", delta.get("metrics", {}))]
        groups.extend((TYPE_LABELS.get(kind, kind), metrics)
                      for kind, metrics in delta.get("by_question_type", {}).items())
        for kind, metrics_by_name in groups:
            for metric, stats in metrics_by_name.items():
                ci95 = stats.get("ci95", [None, None])
                paired_rows.append([
                    delta.get("name"), delta.get("baseline"), kind, METRIC_LABELS.get(metric, metric),
                    stats.get("n", 0), _percent(stats.get("mean")), _percent(stats.get("std")),
                    _percent(ci95[0] if ci95 else None), _percent(ci95[1] if len(ci95) > 1 else None),
                    stats.get("ci_method"),
                ])
    _add_sheet(workbook, "配对比较",
               ["微调结果", "基线结果", "题型", "指标", "有效配对数", "平均差值（百分点）",
                "差值标准差（百分点）", "95% CI下限", "95% CI上限", "区间方法"],
               paired_rows)
    evidence_sheet = _add_sheet(workbook, "检索证据",
                                ["评估名称", "题目ID", "序号", "证据文本", "相关度/距离", "证据完整数据 JSON"],
                                evidence_rows)
    for row in evidence_sheet.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    _add_sheet(workbook, "原始记录", ["评估名称", "题目ID", "分段", "总分段", "完整逐题记录 JSON"], raw_rows)

    info_rows = []
    for key, value in manifest.items():
        _append_chunked_json(info_rows, "本次评估", task.get("name", "evaluation_metrics"), key, value)
    for summary in summaries:
        for key, value in summary.items():
            _append_chunked_json(info_rows, "实验结果", summary.get("name", ""), key, value)
    for key, value in report.items():
        _append_chunked_json(info_rows, "套件对比", "comparison", key, value)
    _add_sheet(workbook, "运行信息", ["类别", "对象", "字段", "值 JSON片段", "分段", "总分段"], info_rows)
    _add_sheet(workbook, "运行日志", ["运行日志"], [[line] for line in log_text.splitlines()])

    cache_rows = _judge_cache_rows(cache_dir)
    if cache_rows:
        cache_sheet = _add_sheet(workbook, "_judge_cache", ["相对路径", "分段", "缓存 JSON片段"], cache_rows)
        cache_sheet.sheet_state = "hidden"

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(dir=workbook_path.parent, prefix=f".{workbook_path.stem}.",
                                         suffix=".xlsx", delete=False) as stream:
            temporary_path = Path(stream.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
    finally:
        workbook.close()
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def _load_config(path: Path) -> list[dict]:
    raw = json.loads(path.read_text(encoding="utf-8-sig"))
    tasks = [raw] if isinstance(raw, dict) else raw
    if not isinstance(tasks, list) or not tasks or not all(isinstance(task, dict) for task in tasks):
        raise ValueError("评估配置必须为对象或非空对象数组。")
    return tasks


def _run_task(task: dict, config_path: Path, cli_generations: list[str] | None,
              cli_output_dir: str | None, cli_resume: bool | None,
              cli_dataset_path: str | None = None) -> int:
    config_dir = config_path.parent
    generation_values = cli_generations if cli_generations is not None else task.get("generations")
    if not isinstance(generation_values, list) or not generation_values:
        raise ValueError("请通过 JSON generations 或 --generations 提供至少一个文件/目录。")
    path_labels = task.get("labels", {})
    if not isinstance(path_labels, dict):
        raise ValueError("labels 必须是从路径到标签的 JSON 对象。")
    entries, labels_seen = [], set()
    for value in generation_values:
        paths = _resolve_generations(value, config_dir)
        for path in paths:
            records = _read_jsonl(path)
            config_names = {row.get("config_name") for row in records if row.get("config_name") is not None}
            if len(config_names) > 1:
                raise ValueError(f"同一 generations 文件含多个 config_name：{path}: {sorted(config_names)}")
            fallback_name = next(iter(config_names), path.parent.name)
            label = path_labels.get(str(value), path_labels.get(str(path), fallback_name))
            label = _safe_label(label)
            if label.casefold() in labels_seen:
                raise ValueError(f"输入标签重复（Windows 不区分大小写）：{label}")
            labels_seen.add(label.casefold())
            entries.append((label, path, records))
    sidecars = [value for _, path, _ in entries if (value := _read_sidecar_evaluation(path)) is not None]
    selected_eval = task.get("evaluation")
    # 兼容将指标开关写在任务顶层的配置形式；推荐写在 evaluation.metrics 中。
    if task.get("metrics") is not None:
        selected_eval = dict(selected_eval or {})
        selected_eval["metrics"] = task["metrics"]
    if selected_eval is None and sidecars:
        selected_eval = sidecars[0]
    evaluation = _normalize_evaluation(selected_eval, config_dir)
    # 输入来自相同实验协议时评价配置应一致；从 sidecar 推断时同样逐一核对。
    normalized_sidecars = [_normalize_evaluation(value, config_dir) for value in sidecars]
    # 指标开关属于本次独立评分配置，不要求与生成阶段 sidecar 的评分选择完全相同。
    # 其余协议参数仍必须一致，避免混合不同题型/裁判口径的结果。
    comparable_evaluation = {key: value for key, value in evaluation.items()
                             if key not in {"metrics", "_metrics_explicit"}}
    comparable_sidecars = [{key: value for key, value in value.items()
                            if key not in {"metrics", "_metrics_explicit"}}
                           for value in normalized_sidecars]
    if any(_canonical(value) != _canonical(comparable_evaluation) for value in comparable_sidecars):
        raise ValueError("所有 generations 对应 summary.json 的 evaluation 设置必须一致，且与评估配置相同。")
    seed = task.get("seed", 42)
    if type(seed) is not int or seed < 0 or seed >= 2**32:
        raise ValueError("seed 必须是 [0, 2**32) 内的整数。")
    output_value = cli_output_dir or task.get("outputDir", "../output/eval_results_metrics")
    output = Path(output_value).expanduser()
    output = (config_dir / output).resolve() if not output.is_absolute() else output.resolve()
    task_count = len(_load_config(config_path))
    task_name = _safe_label(task.get("name") or (f"task_{int(time.time())}" if task_count > 1 else "evaluation_metrics"))
    if task_count > 1 and not cli_output_dir:
        output = output / task_name
    output.mkdir(parents=True, exist_ok=True)
    workbook_path = output / f"{task_name}.xlsx"
    log_buffer = io.StringIO()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s",
                        handlers=[logging.StreamHandler(), logging.StreamHandler(log_buffer)], force=True)
    resume = cli_resume if cli_resume is not None else task.get("resume", True)
    if type(resume) is not bool:
        raise ValueError("resume 必须为布尔值。")
    cache_temporary = tempfile.TemporaryDirectory(prefix="rag-metrics-judge-cache-")
    cache_dir = Path(cache_temporary.name)
    active_metrics = _metric_names(evaluation)
    needs_judge = bool({"factscore", "semantic_score"} & set(active_metrics))
    judge = None
    judge_identity = None
    summaries, by_label = [], {}
    try:
        restored = _restore_judge_cache(workbook_path, cache_dir) if resume else 0
        if restored:
            LOGGER.info("已从现有 Excel 工作簿恢复 %d 个裁判缓存文件。", restored)
        elif not resume:
            LOGGER.info("resume=false：本轮不读取现有工作簿中的裁判缓存。")
        if needs_judge and any(
                any((metric == "factscore" and question_type_for_sample(row) == "open_ended")
                    or (metric == "semantic_score" and
                        (evaluation["judgeScope"] == "all" or
                         evaluation["judgeScope"] == "open_ended" and question_type_for_sample(row) == "open_ended"))
                    for row in records for metric in active_metrics)
                for _, _, records in entries):
            judge = _TextJudgeModel(evaluation["judgeModel"], evaluation["judgeModelPath"])
            judge_identity = {"model_type": evaluation["judgeModel"], "model_path": evaluation["judgeModelPath"],
                              "revision": "local", "software": {name: importlib.metadata.version(name)
                              for name in ("transformers", "torch") if _has_package(name)}}
        persisted_evaluation = {key: value for key, value in evaluation.items() if key != "_metrics_explicit"}
        for label, path, records in entries:
            # JSONL 通常自带 answer/options；若 datasetPath 给出则仅缺字段时回源。
            dataset_path = cli_dataset_path if cli_dataset_path is not None else task.get("datasetPath")
            if dataset_path is not None:
                candidate = Path(dataset_path).expanduser()
                dataset_path = str((config_dir / candidate).resolve() if not candidate.is_absolute() else candidate.resolve())
                records = _load_dataset_fallback(dataset_path, [path], records)
            sidecar = json.loads((path.parent / "summary.json").read_text(encoding="utf-8-sig")) \
                if (path.parent / "summary.json").is_file() else {}
            summary, rows = _save_one(label, path, records, evaluation, seed, judge,
                                       judge_identity, cache_dir, sidecar)
            summaries.append(summary)
            by_label[label] = rows
        report = _write_reports(output, summaries, by_label, evaluation, seed)
        manifest = {
            "name": task.get("name"), "generations": [str(path) for _, path, _ in entries],
            "labels": [label for label, _, _ in entries], "evaluation": persisted_evaluation,
            "seed": seed, "resume": resume, "status": "needs_review" if any(s["judge_failures"] for s in summaries) else "complete",
            "report_status": "complete", "summaries_count": len(summaries), "same_test_data": report["same_test_data"],
            "output_file": str(workbook_path),
        }
        LOGGER.info("正在写入 Excel 评估结果：%s", workbook_path)
        _write_workbook(workbook_path, task, summaries, by_label, evaluation, seed, report, manifest,
                        log_buffer.getvalue(), cache_dir)
        table = _console_table(summaries, evaluation)
        if table:
            print("\n评估结果汇总：")
            print(table)
        print(f"评估完成：{workbook_path}")
        return 2 if any(summary["judge_failures"] for summary in summaries) else 0
    finally:
        if judge is not None:
            judge.close()
        for handler in list(logging.getLogger().handlers):
            logging.getLogger().removeHandler(handler)
            handler.close()
        cache_temporary.cleanup()


def _has_package(name: str) -> bool:
    try:
        importlib.metadata.version(name)
        return True
    except importlib.metadata.PackageNotFoundError:
        return False


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="只评分已有 generations.jsonl，不运行推理。")
    parser.add_argument("--config", default=str(DEFAULT_CONFIG), help="评估 JSON 配置路径。")
    parser.add_argument("--generations", nargs="+", help="generations.jsonl 文件或包含该文件的目录。")
    parser.add_argument("--output_dir", help="覆盖 JSON 中 outputDir；评估结果以同名 Excel 文件保存到此目录。")
    parser.add_argument("--dataset_path", help="缺少答案/选项元数据时回源数据集的路径。")
    parser.add_argument("--resume", dest="resume", action="store_true", default=None,
                        help="复用现有 Excel 工作簿中的裁判缓存。")
    parser.add_argument("--no_resume", dest="resume", action="store_false", help="本轮不读取现有 Excel 工作簿中的裁判缓存。")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        config_path = Path(args.config).expanduser().resolve()
        tasks = _load_config(config_path)
        statuses = [_run_task(task, config_path, args.generations, args.output_dir, args.resume,
                              args.dataset_path) for task in tasks]
        return max(statuses, default=0)
    except Exception as exc:
        print(f"配置或评估失败：{exc}", file=sys.stderr)
        LOGGER.exception("评估失败")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
